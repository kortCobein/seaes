import json, re, hashlib, zipfile, collections, warnings
from pathlib import Path
import openpyxl
warnings.simplefilter('ignore')
BASE=Path(r'C:\Users\cobei\Downloads')
OUT=Path(r'C:\Users\cobei\Desktop\brenda\work')
FILES={'input':next(BASE.glob('SEAES - CAI 2026*.xlsx')),'output':BASE/'SEAES_UTSJR_PLANTILLA_REAL.xlsx','demo':BASE/'SEAES_UTSJR_FICTICIO_ENTREGABLE.xlsx'}
def color(c):
    return {'type':c.type,'value':c.value}
def cell(c):
    return {'ref':c.coordinate,'value':c.value,'style':c.style_id,'kind':c.data_type,'locked':c.protection.locked,'format':c.number_format}
result={}
books={}
for role,path in FILES.items():
    wb=openpyxl.load_workbook(path,data_only=False)
    books[role]=wb
    info={'filename':path.name,'sha256':hashlib.sha256(path.read_bytes()).hexdigest(),'sheetOrder':wb.sheetnames,'definedNames':{n:str(v) for n,v in wb.defined_names.items()},'styles':[], 'sheets':[]}
    for i,s in enumerate(wb._cell_styles):
        fill=wb._fills[s.fillId]
        info['styles'].append({'id':i,'fillId':s.fillId,'fillType':fill.patternType,'fgColor':color(fill.fgColor),'bgColor':color(fill.bgColor),'fontId':s.fontId,'borderId':s.borderId,'numberFormatId':s.numFmtId,'protectionId':s.protectionId})
    for ws in wb:
        all_cells=[c for c in ws._cells.values() if isinstance(c,openpyxl.cell.cell.Cell)]
        nonempty=[cell(c) for c in all_cells if c.value is not None]
        editable=[cell(c) for c in all_cells if not c.protection.locked]
        formulas=[cell(c) for c in all_cells if c.data_type=='f']
        comments=[{'ref':c.coordinate,'text':c.comment.text} for c in all_cells if c.comment]
        validations=[{'sqref':str(v.sqref),'type':v.type,'operator':v.operator,'formula1':v.formula1,'formula2':v.formula2,'allowBlank':v.allowBlank,'error':v.error,'prompt':v.prompt} for v in ws.data_validations.dataValidation]
        charts=[]
        for chart in ws._charts:
            refs=[]
            for ser in chart.series:
                for prop in ['val','cat','xVal','yVal']:
                    elem=getattr(ser,prop,None)
                    if elem:
                        for typ in ['numRef','strRef']:
                            ref=getattr(elem,typ,None)
                            if ref: refs.append(ref.f)
            charts.append({'type':type(chart).__name__,'references':refs})
        entry={'name':ws.title,'state':ws.sheet_state,'dimension':ws.calculate_dimension(),'maxRow':ws.max_row,'maxColumn':ws.max_column,'protection':ws.protection.sheet,'merged':sorted(str(r) for r in ws.merged_cells.ranges),'formulas':formulas,'editableCells':editable,'validations':validations,'nonemptyCells':nonempty,'comments':comments,'charts':charts,'tables':list(ws.tables),'autoFilter':str(ws.auto_filter.ref),'printArea':str(ws.print_area),'stylesUsed':dict(collections.Counter(c.style_id for c in all_cells)),'filledStylesUsed':dict(collections.Counter(c.style_id for c in all_cells if c.value is not None))}
        info['sheets'].append(entry)
    result[role]=info
    print(role,path.name,len(info['sheets']),'sheets',sum(len(s['editableCells']) for s in info['sheets']),'unlocked',sum(len(s['formulas']) for s in info['sheets']),'formulas')
diffs=[]
for ws in books['output']:
    if ws.title not in books['demo']: continue
    other=books['demo'][ws.title]
    changed=[]
    for row in range(1,max(ws.max_row,other.max_row)+1):
        for col in range(1,max(ws.max_column,other.max_column)+1):
            a,b=ws.cell(row,col),other.cell(row,col)
            if a.value!=b.value: changed.append({'ref':a.coordinate,'from':a.value,'to':b.value,'style':a.style_id,'locked':a.protection.locked,'kind':b.data_type})
    diffs.append({'sheet':ws.title,'changes':changed})
result['outputDemoDifferences']=diffs
(OUT/'excel-audit.json').write_text(json.dumps(result,ensure_ascii=False,indent=2,default=str),encoding='utf-8')
summary=[]
for role in ['input','output','demo']:
    summary.append('\n'+role.upper())
    for s in result[role]['sheets']:
        summary.append(f"{s['name']} | {s['dimension']} | {len(s['nonemptyCells'])} filled | {len(s['editableCells'])} unlocked | {len(s['formulas'])} formulas | {len(s['validations'])} validations | {len(s['charts'])} charts | protected {s['protection']}")
(OUT/'excel-overview.txt').write_text('\n'.join(summary),encoding='utf-8')
print('\n'.join(summary))
